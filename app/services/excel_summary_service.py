"""Excel 单元格定向汇总服务。"""

from __future__ import annotations

import gc
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Any
import warnings

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl.styles.stylesheet")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.rule import Rule, validate_rules
from app.services.file_filter_service import list_excel_files


LogCallback = Callable[[str, str], None]
ProgressCallback = Callable[[int, int], None]
CurrentFileCallback = Callable[[str], None]


@dataclass
class SummaryResult:
    """汇总执行结果。"""

    output_path: Path
    total_files: int
    processed_files: int


def normalize_output_path(output_file: str | Path) -> Path:
    """规范输出路径，无后缀时自动补 .xlsx。"""
    output_path = Path(output_file).expanduser()
    if not output_path.suffix:
        output_path = output_path.with_suffix(".xlsx")
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    return output_path


def run_summary(
    target_folder: str | Path,
    output_file: str | Path,
    keyword: str,
    filter_mode: str,
    rules: list[Rule],
    log_callback: LogCallback | None = None,
    progress_callback: ProgressCallback | None = None,
    current_file_callback: CurrentFileCallback | None = None,
) -> SummaryResult:
    """执行 Excel 单元格定向汇总。"""
    normalized_rules = validate_rules(rules)
    output_path = normalize_output_path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _ensure_output_writable(output_path)

    files = list_excel_files(target_folder, keyword=keyword, filter_mode=filter_mode)
    total = len(files)
    _log(log_callback, "INFO", f"找到 {total} 个待处理 Excel 文件。")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "汇总结果"
    headers = ["文件名", "文件路径", *[rule.output_column for rule in normalized_rules]]
    worksheet.append(headers)
    _apply_header_style(worksheet)

    processed = 0
    for file_path in files:
        processed += 1
        if current_file_callback is not None:
            current_file_callback(str(file_path))
        _log(log_callback, "INFO", f"正在处理：{file_path.name}")
        row_values = [file_path.name, str(file_path)]
        row_values.extend(_read_file_values(file_path, normalized_rules, log_callback))
        worksheet.append(row_values)
        if progress_callback is not None:
            progress_callback(processed, total)

    _adjust_columns(worksheet)
    workbook.save(output_path)
    workbook.close()
    _log(log_callback, "DONE", f"汇总完成，输出文件：{output_path}")

    del workbook, files, normalized_rules
    gc.collect()
    return SummaryResult(output_path=output_path, total_files=total, processed_files=processed)


def _read_file_values(
    file_path: Path,
    rules: list[Rule],
    log_callback: LogCallback | None,
) -> list[Any]:
    values: list[Any] = []
    workbook = None
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        for rule in rules:
            values.append(_read_rule_value(workbook, rule, file_path, log_callback))
    except PermissionError as exc:
        _log(log_callback, "ERROR", f"读取失败，文件可能被占用：{file_path.name}；{exc}")
        values = [None for _ in rules]
    except Exception as exc:
        _log(log_callback, "ERROR", f"读取失败：{file_path.name}；{exc}")
        values = [None for _ in rules]
    finally:
        if workbook is not None:
            workbook.close()
        gc.collect()
    return values


def _read_rule_value(
    workbook: Any,
    rule: Rule,
    file_path: Path,
    log_callback: LogCallback | None,
) -> Any:
    sheet_name = _locate_sheet(workbook.sheetnames, rule)
    if sheet_name is None:
        _log(
            log_callback,
            "WARN",
            f"{file_path.name} 未找到 Sheet：模式={rule.sheet_mode}，值={rule.sheet_value}",
        )
        return None
    try:
        worksheet = workbook[sheet_name]
        return worksheet[rule.cell].value
    except Exception as exc:
        _log(log_callback, "WARN", f"{file_path.name} 读取单元格 {rule.cell} 失败：{exc}")
        return None


def _locate_sheet(sheet_names: list[str], rule: Rule) -> str | None:
    if rule.sheet_mode == "exact":
        return rule.sheet_value if rule.sheet_value in sheet_names else None
    if rule.sheet_mode == "contains":
        return next((name for name in sheet_names if rule.sheet_value in name), None)
    if rule.sheet_mode == "index":
        index = int(rule.sheet_value) - 1
        return sheet_names[index] if 0 <= index < len(sheet_names) else None
    return None


def _ensure_output_writable(output_path: Path) -> None:
    if output_path.exists() and output_path.name.startswith("~$"):
        raise PermissionError("输出文件是 Excel 临时文件，请选择正常 .xlsx 文件。")
    try:
        with output_path.open("a+b"):
            pass
    except PermissionError as exc:
        raise PermissionError("输出文件可能正在被 Excel 占用，请关闭后重试。") from exc


def _apply_header_style(worksheet: Any) -> None:
    fill = PatternFill("solid", fgColor="E8EEF7")
    font = Font(bold=True, color="1F2937")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"


def _adjust_columns(worksheet: Any) -> None:
    for column_cells in worksheet.columns:
        max_length = 10
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value) + 2, 48))
        worksheet.column_dimensions[get_column_letter(column)].width = max_length


def _log(log_callback: LogCallback | None, level: str, message: str) -> None:
    if log_callback is not None:
        log_callback(level, message)
